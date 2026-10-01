from pathlib import Path

import pytest

from src.ingest import jerarquia as jq
from src.ingest.conversores._texto import Parrafo, segmentar
from src.ingest.conversores.base import Bloque
from src.ingest.conversores.html_plantilla import HtmlPlantilla
from src.ingest.markdown import a_markdown, niveles
from src.ingest.normalizar import desguionar, texto


@pytest.mark.parametrize("linea,num,resto", [
    ("ARTICULO 1o.", "1", ""),
    ("ARTÍCULO 12-A. EPÍGRAFE", "12A", "EPÍGRAFE"),
    ("ARTÍCULO 240-1. OTRAS RENTAS", "240-1", "OTRAS RENTAS"),
    ("Artículo 5 bis. Texto", "5-bis", "Texto"),
    ("ARTÍCULO 2.2.1.1.1. Objeto", "2.2.1.1.1", "Objeto"),
    ("ARTICULO TRANSITORIO 1. Texto", "transitorio-1", "Texto"),
    ("ARTÍCULO  1º. El artículo 23", "1", "El artículo 23"),
    ("ARTICULO ÚNICO. Texto", "unico", "Texto"),
])
def test_articulo(linea, num, resto):
    assert jq.articulo(linea) == (num, resto)


@pytest.mark.parametrize("linea", ["Artículos 3 y 4 de la ley", "El artículo 5", "ARTICULADO"])
def test_no_es_articulo(linea):
    assert jq.articulo(linea) is None


@pytest.mark.parametrize("linea,tipo,nombre", [
    ("TITULO I.", "titulo", ""),
    ("TÍTULO PRELIMINAR.", "titulo", ""),
    ("TÍTULO TRANSITORIO.", "titulo", ""),
    ("PRIMERA PARTE.", "parte", ""),
    ("LIBRO PRIMERO", "libro", ""),
    ("CAPITULO X. DESCUENTOS TRIBUTARIOS", "capitulo", "DESCUENTOS TRIBUTARIOS"),
    ("SECCIÓN 2a", "seccion", ""),
])
def test_estructura(linea, tipo, nombre):
    assert jq.estructura(linea) == (tipo, nombre)


@pytest.mark.parametrize("linea", ["Título valor de algo", "Parte de la doctrina", "Capítulo aparte"])
def test_no_es_estructura(linea):
    assert jq.estructura(linea) is None


@pytest.mark.parametrize("linea,es", [
    ("I. ANTECEDENTES", True),
    ("II . Consideraciones y fundamentos", True),
    ("VI. CONSIDERACIONES Y FUNDAMENTOS DE LA CORTE", True),
    ("RESUELVE:", True),
    ("SALVAMENTO DE VOTO DEL MAGISTRADO X", True),
    ("I. ANTECEDENTES (3)", False),  # entrada del índice
    ("III. Los artículos impugnados vulneran la igualdad", False),
    ("COMPAÑÍA DE SEGUROS S.A.", False),
])
def test_seccion_sentencia(linea, es):
    assert bool(jq.seccion_sentencia(linea)) is es


def test_normalizar():
    assert texto("modiﬁcado  por­la  ley") == "modificado porla ley"
    assert desguionar(["el contra-", "to de trabajo"]) == "el contrato de trabajo"
    assert desguionar(["Decreto-", "Ley 2351"]) == "Decreto- Ley 2351"


def test_niveles_dinamicos_sin_huecos():
    bloques = [Bloque("capitulo", "CAPÍTULO I."), Bloque(jq.ARTICULO, "Artículo 1.", "1")]
    assert niveles(bloques) == {"capitulo": 2, jq.ARTICULO: 3}


def test_segmentar_descarta_articulos_citados():
    parrafos = [Parrafo("ARTÍCULO 48. El artículo 369 del Código quedará así:"),
                Parrafo("ARTÍCULO 369. Modificación de los Estatutos. Toda modificación…"),
                Parrafo("Artículo 23. Elementos esenciales."),
                Parrafo("ARTÍCULO 49. Texto.")]
    bloques, stats = segmentar(parrafos, sentencia=False)
    assert [b.etiqueta for b in bloques if b.tipo == jq.ARTICULO] == ["48", "49"]
    assert stats["articulos_rechazados_por_secuencia"] == 1


def test_segmentar_parte_articulo_pegado():
    bloques, _ = segmentar([Parrafo("ARTÍCULO 59. Texto en lo pertinente. ARTICULO 60. El artículo 434.")], False)
    assert [b.etiqueta for b in bloques if b.tipo == jq.ARTICULO] == ["59", "60"]


PAGINA = """<html><body><div id="selector_aj"><select><option>menu</option></select></div>
<!--Inicio documento-->
<p class="centrado"><a class="bookmarkaj" name="Nivel001">TITULO I. </A></p>
<p class="centrado"><span class="b_aj">DE LOS PRINCIPIOS</span></p>
<p><a class="bookmarkaj" name="1">ARTICULO 1o.</A> Colombia es un <S>Estado</S> social.</p>
<div><a class="caja_vja_encabezado" href="javascript:insRow1()">Concordancias</a></div>
<table id="Table1"><tbody><tr><td><p>Ley 388 de 1997</p></td></tr></tbody></table>
<p class="centrado"><A name="Nivel002"></A></p>
<p class="centrado"><span class="b_aj">CAPITULO 1.</span></p>
<p class="centrado">DE LOS DERECHOS</p>
<p><span><A name="2">ARTICULO 2o. FINES.</A></span> Son fines.</p>
<p class="centrado">TÍTULO I OBJETO</p>
<!--Fin documento-->
<div id="logo_aj">Derechos de autor reservados</div></body></html>"""


def test_plantilla_senado(tmp_path: Path):
    f = tmp_path / "ley_p000.html"
    f.write_text(PAGINA, encoding="cp1252")
    res = HtmlPlantilla().convertir([f], sentencia=False)
    tipos = [(b.tipo, b.texto) for b in res.bloques]
    assert tipos[0] == ("titulo", "TITULO I. DE LOS PRINCIPIOS")
    assert tipos[1] == (jq.ARTICULO, "Artículo 1.")
    assert tipos[2] == ("parrafo", "Colombia es un ~~Estado~~ social.")
    assert tipos[3] == ("caja", "Ley 388 de 1997")
    assert tipos[4] == ("capitulo", "CAPITULO 1. DE LOS DERECHOS")
    assert tipos[5] == (jq.ARTICULO, "Artículo 2. FINES")
    # Encabezado sin ancla después del último artículo: texto anexo, no jerarquía.
    assert tipos[-1] == ("parrafo", "TÍTULO I OBJETO")
    assert res.stats["articulos_fuente"] == 2

    md = a_markdown({"doc_id": "ley"}, "Ley", res.bloques)
    assert "## TITULO I. DE LOS PRINCIPIOS" in md
    assert "### CAPITULO 1. DE LOS DERECHOS" in md
    assert "#### Artículo 2. FINES" in md
    assert "> **Concordancias:**" in md
    assert "menu" not in md and "Derechos de autor" not in md


# --- Formatos y metadatos agregados al convertir el corpus completo -----------------

def test_leer_html_mezcla_utf8_y_cp1252(tmp_path: Path):
    from src.ingest.conversores._html import leer
    f = tmp_path / "mixto.htm"
    # "artículo" en UTF-8 y "país" en windows-1252 en el mismo archivo, con charset declarado.
    f.write_bytes('<meta charset="windows-1252"><p>artículo '.encode("utf-8") + "país</p>".encode("cp1252"))
    t = leer(f)
    assert "artículo" in t and "país" in t and "charset" not in t


def test_segmentar_numeracion_decimal_ignora_articulo_citado():
    ps = [Parrafo(t) for t in ("Artículo 1.2.1.22.5. Texto", "Artículo 206E. Citado del Estatuto",
                               "Artículo 1.2.1.22.6. Texto", "Artículo 1.311.8.1.3. Número mal extraído",
                               "Artículo 1.2.1.23.1. Texto")]
    bloques, _ = segmentar(ps, sentencia=False)
    assert [b.etiqueta for b in bloques if b.tipo == jq.ARTICULO] == ["1.2.1.22.5", "1.2.1.22.6", "1.2.1.23.1"]


def test_tabla_markdown():
    from src.ingest.conversores.base import TABLA
    md = a_markdown({"doc_id": "x"}, "X", [Bloque(TABLA, "", filas=[["Código", "Tarifa"], ["8462", "5 %"], ["a|b"]])])
    assert "| Código | Tarifa |\n|---|---|\n| 8462 | 5 % |\n" + r"| a\|b |   |" in md


def test_clave_encabezado_ignora_orden_y_numeros():
    from src.ingest.conversores.pdf import _clave, _ruido
    assert _clave("107 del Hoja No. 2 .. Decreto No.") == _clave("Hoja No. 4 del Decreto No. 107")
    assert _ruido("REPÚBLICA DE COLOMBIA .. t~SIOtNLII m e· :: ,·utsuG ~ SECRfTAH//.'")
    assert not _ruido("ARTÍCULO 2. Vigencia. El presente decreto rige a partir de su publicación.")


@pytest.mark.parametrize("did,esperado", [
    ("ley_1581_2012", {"tipo_norma": "ley", "numero": "1581", "anio": "2012", "organo_emisor": "Congreso de la República"}),
    ("codigo_general_proceso", {"tipo_norma": "ley", "numero": "1564", "nombre_citable": "Código General del Proceso"}),
    ("estatuto_tributario", {"tipo_norma": "decreto", "numero": "624", "anio": "1989"}),
    ("jurisprudencia_c-355_2006", {"numero": "C-355", "organo_emisor": "Corte Constitucional"}),
    ("jurisprudencia_sl-4283_2021", {"organo_emisor": "Corte Suprema de Justicia", "sala": "Sala de Casación Laboral"}),
    ("jurisprudencia_ce-05001-23-26-000-1994-02321-01_2012", {"organo_emisor": "Consejo de Estado", "anio": "2012"}),
])
def test_metadatos(did, esperado):
    from src.ingest.metadatos import metadatos
    m = metadatos(did, {}, {"vigencia": None})
    assert {k: m.get(k) for k in esperado} == esperado and m["vigencia"] == "sin_marca"


def test_docx(tmp_path: Path):
    import docx
    from src.ingest.conversores.docx import Docx
    d = docx.Document()
    d.add_paragraph("I. ANTECEDENTES")
    p = d.add_paragraph("Texto vigente y ")
    p.add_run("texto declarado inexequible").font.strike = True
    t = d.add_table(rows=2, cols=2)
    for i, fila in enumerate([["Año", "Valor"], ["2024", "47.065"]]):
        for j, v in enumerate(fila):
            t.cell(i, j).text = v
    f = tmp_path / "s.docx"
    d.save(f)
    res = Docx().convertir([f], sentencia=True)
    md = a_markdown({"doc_id": "s"}, "S", res.bloques)
    assert "## I. ANTECEDENTES" in md or "ANTECEDENTES" in md
    assert "~~texto declarado inexequible~~" in md
    assert "| Año | Valor |" in md


def test_unir_spans_separa_palabras_de_capas_ocr():
    from src.ingest.conversores.pdf import _unir_spans
    spans = [{"text": "audiencia", "bbox": (10, 0, 60, 10), "size": 10},
             {"text": "de", "bbox": (63, 0, 73, 10), "size": 10},
             {"text": "acusación", "bbox": (76, 0, 120, 10), "size": 10},
             {"text": ",", "bbox": (120.2, 0, 122, 10), "size": 10}]
    assert _unir_spans(spans) == "audiencia de acusación,"


def test_capa_mala_de_ocr_ajeno():
    from src.ingest.conversores.pdf import Linea, _capa_mala
    mala = [Linea(0, 0, 1, "Rep正博ぐadeCoIomhia Corte Suprema de Justicia", False)]
    buena = [Linea(0, 0, 1, "En el mes de mayo de 2011, en horas de la noche y cuando se hallaba dedicado " * 3, False)]
    assert _capa_mala(mala) and not _capa_mala(buena)


# --- Recuperación: tokenización BM25, fusión y chunking --------------------------------

def test_tokens_bm25_identificadores_y_tildes():
    from src.knowledge.tokenization import tokens
    t = tokens("Sentencia C-355 de 2006, artículo 240-1 del Decreto 1.625; artículo 5")
    assert "c_355" in t and "240_1" in t and "1625" in t and "5" in t and "2006" in t
    assert tokens("C-748-11") == tokens("C-748 de 2011") == tokens("c-748/11") == ["c_748", "2011"]
    assert tokens("T-025 de 2004")[:1] == ["t_25"] and tokens("SU-214/16") == ["su_214", "2016"]
    assert tokens("CONSTITUCIÓN") == tokens("constitucion")


def test_rrf_desempata_por_id():
    from src.knowledge.hybrid_search import rrf
    assert [c for c, _ in rrf([["b", "a"], ["a", "b"]])] == ["a", "b"]
    assert rrf([["x", "y"], ["y"]])[0][0] == "y"


def test_chunker_articulo_con_ruta_y_notas(tmp_path: Path):
    from src.knowledge.chunking import chunks_de
    md = tmp_path / "ley_1_2000.md"
    md.write_text('---\ndoc_id: "ley_1_2000"\ntipo_documento: "norma"\ntipo_norma: "ley"\nnumero: "1"\n'
                  'anio: "2000"\nnombre_citable: "Ley 1 de 2000"\n---\n\n# Ley 1 de 2000\n\nPor la cual...\n\n'
                  '## TÍTULO I. DISPOSICIONES\n\n### Artículo 1. OBJETO\n\nTexto del artículo.\n\n'
                  '> **Notas de Vigencia:**\n>\n> Modificado por la Ley 2 de 2001.\n\n'
                  '> **Legislación Anterior:**\n>\n> Texto viejo.\n', encoding="utf-8")
    cs = {c["chunk_id"]: c for c in chunks_de(md)}
    art = cs["ley_1_2000/art_1"]
    assert art["texto"].startswith("Ley 1 de 2000 › TÍTULO I. DISPOSICIONES\nArtículo 1. OBJETO")
    assert "Modificado por la Ley 2 de 2001" in art["texto"] and "Texto viejo" not in art["texto"]
    assert "Texto viejo" in cs["ley_1_2000/art_1/notas"]["texto"]
    assert "ley_1_2000/preambulo#1" in cs


def test_seleccion_tope_sentencias_y_penalizacion():
    from src.knowledge.hybrid_search import Config, Recuperador
    rec = Recuperador.__new__(Recuperador)
    rec.cfg = Config(k=4, max_sentencias=2, max_por_doc=3, dedup=0.0)  # textos de ejemplo iguales
    datos = {f"s{i}/ficha": {"doc_id": f"jurisprudencia_t-{i}_2020", "tipo_chunk": "ficha",
                             "tipo_documento": "sentencia", "texto": "t"} for i in range(5)}
    datos["ley_1_2000/art_1"] = {"doc_id": "ley_1_2000", "tipo_chunk": "articulo", "tipo_documento": "norma",
                                 "articulo_id": "ley_1_2000/art_1", "texto": "t"}
    orden = [(f"s{i}/ficha", 0.9 - i / 100) for i in range(5)] + [("ley_1_2000/art_1", 0.5)]
    ids = [p["chunk_id"] for p in rec._seleccionar(orden, datos)]
    # Tope de 2 sentencias: entra la norma; la 3.ª ficha vuelve solo para completar k=4. Salen por puntaje.
    assert ids == ["s0/ficha", "s1/ficha", "s2/ficha", "ley_1_2000/art_1"]
    rec.cfg = Config(k=3, max_sentencias=2, max_por_doc=3, dedup=0.0)
    assert [p["chunk_id"] for p in rec._seleccionar(orden, datos)] == ["s0/ficha", "s1/ficha", "ley_1_2000/art_1"]
